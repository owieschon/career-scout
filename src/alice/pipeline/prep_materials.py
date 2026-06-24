"""Application materials generation — Alice's prep workflow.

Reads:  feedback/prep-queue.json (pending prep requests from Jordan Avery)
        feedback/question-answers.jsonl (Jordan Avery's answers to targeted questions)
        applications/<slug>/ (existing draft state per role)
Writes: applications/<slug>/{resume-draft,cover-letter-draft,targeted-questions,application-strategy,.metadata}.md/json
        applications/<slug>/{resume-final,cover-letter-final}.md (after Jordan Avery answers)

Two paths:
  1. NEW prep request -> generate drafts + questions, mark questions pending
  2. New question answers received -> integrate into final drafts
"""
import json
import re
import sys
import urllib.parse
from datetime import datetime
from pathlib import Path

from alice import repo_paths
from alice.llm import llm
from alice.persistence import ledger
from alice import safe_state
from alice.pipeline.source_deep import get as fetch_json, strip as html_strip

QUEUE = Path(str(repo_paths.FEEDBACK / "prep-queue.json"))
ANSWERS = Path(str(repo_paths.FEEDBACK / "question-answers.jsonl"))
APPS_DIR = Path(str(repo_paths.APPLICATIONS))
TEMPLATES = Path(str(repo_paths.TEMPLATES))


def _slug(company, role):
    s = f"{company}-{role}".lower()
    s = re.sub(r"[^a-z0-9-]+", "-", s)
    return re.sub(r"-+", "-", s).strip("-")[:80]


def _read_resume_text(variant_filename):
    """Read a .docx resume variant as plain text (paragraph per line)."""
    path = TEMPLATES / variant_filename
    if not path.exists():
        return ""
    try:
        from docx import Document
        d = Document(str(path))
        return "\n".join(p.text for p in d.paragraphs if p.text.strip())
    except Exception:
        return ""


def _pick_resume_variant(archetype):
    """Pick the right master resume variant based on the role's archetype."""
    mapping = {
        "AE":        "resume-senior-ae.docx",
        "RevOps":    "resume-revenue-architect.docx",
        "TAM / CS":  "resume-tam.docx",
        "FDE":       "resume-operator-builder.docx",
    }
    return mapping.get(archetype, "resume-operator-builder.docx")


def _fetch_jd(url):
    """Best-effort JD body fetch — reuses enrich_hypotheses logic."""
    try:
        from alice.pipeline.enrich_hypotheses import fetch_jd
        return fetch_jd(url)
    except Exception:
        return None


def _row_for_substr(substr):
    """Find one sheet row matching substring; return (row_idx, row_dict) or None."""
    from alice.notify.imap_reply import _match_sheet_row
    ws = ledger._ws()
    rows = ws.get_all_records()
    hits = _match_sheet_row(substr, rows)
    if len(hits) != 1:
        return None
    return hits[0]


def generate_drafts(row_idx, row, rush=False):
    """Generate resume-draft, cover-letter-draft, targeted-questions, application-strategy.
    Returns dict with paths + Alice's metadata."""
    company = row.get("company", "")
    role_title = row.get("role", "")
    url = row.get("url", "")
    comp = row.get("comp", "n/d")
    rationale = row.get("rationale", "")
    archetype_match = re.match(r"\[([^\]]+)\]", rationale)
    archetype = archetype_match.group(1) if archetype_match else "Other"

    slug = _slug(company, role_title)
    pkg_dir = APPS_DIR / slug
    pkg_dir.mkdir(parents=True, exist_ok=True)

    jd_body = _fetch_jd(url) or "(JD body could not be fetched; reasoning from title + rationale only)"
 # JD-body persistence: parallel to prep_pipeline._persist_jd_body — write
 # the raw fetched JD to applications/<slug>/jd-body.txt at fetch time so
 # the body survives URL rot. Non-fatal on failure. Skip the placeholder
 # ("not fetched") string — only real fetches go on disk.
    if jd_body and not jd_body.startswith("(JD body could not"):
        try:
            jd_path = pkg_dir / "jd-body.txt"
            header = (
                f"# JD body persisted by prep_materials at draft time\n"
                f"# url:        {url}\n"
                f"# fetched_at: {datetime.now().isoformat(timespec='seconds')}\n"
                f"# company:    {company}\n"
                f"# role:       {role_title}\n"
                f"# chars:      {len(jd_body)}\n"
                f"# ---\n"
            )
            jd_path.write_text(header + jd_body)
        except Exception as e:
            print(f"  [persist jd-body: {type(e).__name__}: {e}]")
    variant_name = _pick_resume_variant(archetype)
    variant_text = _read_resume_text(variant_name)

    brief = llm.load_alice_brief()

 # ----- 1. Resume draft (Opus per llm.py defaults — voice + strategic framing matter) -----
    resume_prompt = f"""Write a resume DRAFT for Jordan Avery targeted at this role.

ROLE
Company: {company}
Title: {role_title}
URL: {url}
Comp: {comp}
Archetype: {archetype}
Why this role qualified (Alice's earlier read): {rationale}

JD BODY
{jd_body[:6000]}

REFERENCE MASTER VARIANT (what's TRUE about Jordan Avery, organized as he currently presents)
{variant_text}

YOUR JOB
Per your brief (Per-application strategic execution section): write the OPTIMAL NARRATIVE for
this role, not a tailoring-down of the variant. You have creative freedom on framing, ordering,
and emphasis. Where you need a specific candidate-fact that isn't in the variant, leave a
[FILL: <specific question>] placeholder.

CONSTRAINTS
- Plain markdown, no docx structure
- Top section: name + contact + headline (1 line) + summary (3-4 sentences)
- Then experience section: each role with company, title, dates, 3-5 bullets
- Then technical/skills section
- Then education
- Total length: fits one printed page if rendered as resume (~700 words)
- Jordan Avery's voice: direct, evidence-driven, no em dashes, no "passionate", no "leveraged", no "synergies"
- Bullets lead with outcome, not activity
- Numbers and names whenever defensible
"""
    print(f"    generating resume draft (model: {llm.MODEL_FOR_TASK['resume_draft']})...")
    resume_res = llm.call("resume_draft", resume_prompt, system=brief, max_tokens=2500)
    (pkg_dir / "resume-draft.md").write_text(resume_res["text"])

 # ----- 2. Cover letter draft (Opus) -----
    cover_prompt = f"""Write a cover letter DRAFT for Jordan Avery targeted at this role.

ROLE
Company: {company}
Title: {role_title}
URL: {url}
Archetype: {archetype}
Why this role qualified: {rationale}

JD BODY
{jd_body[:6000]}

REFERENCE MASTER VARIANT (Jordan Avery's true experience)
{variant_text[:2000]}

YOUR JOB
Write the OPTIMAL opening. Per your brief: a specific reason for this role specifically,
NOT a generic interest hook. Three to five short paragraphs. Where you need a specific
candidate-fact not in the variant, leave [FILL: <specific question>] placeholders.

CONSTRAINTS
- First person, Jordan Avery's voice
- No em dashes, no "passionate", no consulting-speak
- Opens with specific evidence of why THIS role
- Names 2-3 candidate-evidence pairs that map to JD signals
- Closes with what Jordan Avery would bring in the first 90 days
- Sign off plainly (no "Sincerely yours" theater)
"""
    print(f"    generating cover letter draft (model: {llm.MODEL_FOR_TASK['cover_letter_draft']})...")
    cover_res = llm.call("cover_letter_draft", cover_prompt, system=brief, max_tokens=1500)
    (pkg_dir / "cover-letter-draft.md").write_text(cover_res["text"])

 # ----- 3. Targeted questions (Haiku — short structured output) -----
    questions_prompt = f"""You just wrote a resume draft and cover letter draft for Jordan Avery at {company} ({role_title}).
Both drafts contain [FILL: ...] placeholders where you need specific candidate-facts that weren't in his master variant.

Resume draft:
{resume_res['text'][:3000]}

Cover letter draft:
{cover_res['text'][:1500]}

YOUR JOB
List the 5-12 targeted questions Jordan Avery needs to answer so you can integrate his specifics into the final drafts.
Each question should be narrow, specific, and answerable in 1-2 sentences. Number them.

OUTPUT FORMAT (exactly this)
1. <question>
   placeholder: [FILL: <verbatim placeholder from above>]
2. <question>
   placeholder: [FILL: <...>]
...

Jordan Avery replies with 'answer 1: <text>' through 'answer N: <text>' to fill them.
"""
    print(f"    generating targeted questions (model: {llm.MODEL_FOR_TASK['targeted_questions']})...")
    q_res = llm.call("targeted_questions", questions_prompt, system=brief, max_tokens=1200)
    (pkg_dir / "targeted-questions.md").write_text(q_res["text"])

 # ----- 4. Application strategy (Sonnet) -----
    strategy_prompt = f"""Write the application-strategy note for Jordan Avery's screen prep on this role.

ROLE
Company: {company}
Title: {role_title}
Archetype: {archetype}
JD signal summary: {rationale}

JD BODY
{jd_body[:5000]}

YOUR JOB (internal note for Jordan Avery — not part of the application)
Cover:
1. THE STORY YOU'RE TELLING — the narrative arc the resume + cover both support
2. TRANSFER ANGLES YOU'RE LEANING ON — specific JD-signal-to-candidate-evidence pairs that justify the fit
3. EMPHASIZE AT FIRST SCREEN — the 2-3 things Jordan Avery should make sure to mention if not asked
4. KNOWN RISKS the hiring manager will surface — and how to pre-empt

CONSTRAINTS
- Direct, evidence-driven, no em dashes
- Specific, not generic
- This is Jordan Avery's screen prep, not marketing copy
"""
    print(f"    generating application strategy (model: {llm.MODEL_FOR_TASK['application_strategy']})...")
    s_res = llm.call("application_strategy", strategy_prompt, system=brief, max_tokens=1500)
    (pkg_dir / "application-strategy.md").write_text(s_res["text"])

 # ----- 4.5. Outreach targets (Sonnet) -----
 # Targeting research informs the application itself (cover letter framing,
 # screen prep, whether to apply cold vs wait for a warm intro), so it runs
 # at prep time. Only generated if not already present — protects pre-seeded
 # hand-written intel (e.g. a hand-curated warm path).
    outreach_path = pkg_dir / "outreach-targets.md"
    o_cost = 0.0
    o_model = None
    if outreach_path.exists():
        print(f"    outreach-targets.md exists (pre-seeded, not regenerating)")
    else:
        outreach_prompt = f"""Research outreach targets for Jordan Avery's application to this role.
Identify likely decision-makers and surface warm-path categories Jordan Avery should investigate.

ROLE
Company: {company}
Title: {role_title}
JD signal summary: {rationale}

JD BODY
{jd_body[:5000]}

YOUR JOB
Produce a structured outreach-targets document with these sections:

## DECISION-MAKER TARGETS (public-source-identifiable)
For each: name (if identifiable from JD or commonly-known role at company), title,
why-them (likely hiring manager / likely first-screen interviewer / function owner),
where-found citation (e.g. "company team page", "LinkedIn search recommended for VP CS at <company>"),
and one specific reason this person matters for THIS role.
Be honest when you can't identify by name — say "likely hiring manager is the head of
<function>; identifiable via LinkedIn search at <company>". Don't fabricate names.

## WARM-PATH CATEGORIES TO INVESTIGATE
List the connection categories worth Jordan Avery searching his own network for:
- alma mater (note specific schools that show up in company leadership / engineering)
- prior-employer overlap (companies in <company>'s history Jordan Avery has worked at or near)
- industry mutual-connection candidates (verticals Jordan Avery has credibility in that overlap with this company's customer base)
- mutual investor/board connections if applicable

## CANDIDATE-PROVIDED WARM PATHS
(Leave this section as a labeled placeholder for Jordan Avery to fill in via chat with
personal-network intel: "I know <name>", "<name> is my <relationship>",
"connected via <person> on LinkedIn". Format the placeholder so it's clear Jordan Avery
should write here.)

## RECOMMENDED APPROACH
One paragraph: should Jordan Avery apply cold, pursue warm intro first, or both in parallel?
Justified by the targets + paths above. Honest when the answer is "apply cold, no
warm path obvious yet".

CONSTRAINTS
- Direct, evidence-driven, no em dashes
- Don't fabricate names of specific people — identify by role/title if name not known
- Don't pad with generic outreach advice
"""
        print(f"    generating outreach targets (model: {llm.MODEL_FOR_TASK['application_strategy']})...")
        o_res = llm.call("application_strategy", outreach_prompt, system=brief, max_tokens=1500)
        outreach_path.write_text(o_res["text"])
        o_cost = o_res["cost_usd"]
        o_model = o_res["model"]

 # ----- 5. Metadata -----
    total_cost = resume_res["cost_usd"] + cover_res["cost_usd"] + q_res["cost_usd"] + s_res["cost_usd"] + o_cost
    metadata = {
        "company":          company,
        "role":             role_title,
        "row_idx":          row_idx,
        "archetype":        archetype,
        "url":              url,
        "comp":             comp,
        "rationale":        rationale,
        "generated_at":     datetime.now().isoformat(timespec="seconds"),
        "variant_used":     variant_name,
        "rush":             rush,
        "draft_costs": {
            "resume":   resume_res["cost_usd"],
            "cover":    cover_res["cost_usd"],
            "questions": q_res["cost_usd"],
            "strategy": s_res["cost_usd"],
            "total":    total_cost,
        },
        "models_used": {
            "resume":   resume_res["model"],
            "cover":    cover_res["model"],
            "questions": q_res["model"],
            "strategy": s_res["model"],
        },
        "answers_received": [],
        "final_generated":  False,
    }
    (pkg_dir / ".metadata.json").write_text(json.dumps(metadata, indent=2))

 # ----- 6. Update sheet status to materials pending -----
 # A fresh-auth read-back of column G runs after the gspread write.
 # sheet_write_ok = API call returned without raising.
 # sheet_write_verified = an independent surface confirms the write landed.
    from alice.ops import verify
    sheet_write_ok = False
    sheet_write_verified = False
    sheet_write_error = None
    sheet_verify_claim = None
    ws = ledger._ws()
    try:
        ledger.update_status(
            ws, row_idx, "materials pending",
            source="prep_materials:auto_after_package_build",
        )
        sheet_write_ok = True
        vr = verify.verify_sheet_status_write(row_idx, "materials pending")
        sheet_write_verified = bool(vr.ok)
        sheet_verify_claim = vr.claim
        if vr.ok:
            print(f"    [sheet write 'materials pending' row {row_idx}: VERIFIED via fresh-auth read-back]")
        else:
            print(f"    [sheet write 'materials pending' row {row_idx}: WRITE ATTEMPTED, "
                  f"VERIFY FAILED: {vr.claim}]")
    except Exception as e:
        sheet_write_error = str(e)[:200]
        print(f"    [sheet write 'materials pending' row {row_idx}: FAILED — {sheet_write_error}]")

    return {
        "slug":                  slug,
        "pkg_dir":               str(pkg_dir),
        "total_cost":            total_cost,
        "sheet_write_ok":        sheet_write_ok,
        "sheet_write_verified":  sheet_write_verified,
        "sheet_write_error":     sheet_write_error,
        "sheet_verify_claim":    sheet_verify_claim,
    }


def integrate_answers(slug):
    """When Jordan Avery answers targeted questions, integrate them into resume-final + cover-letter-final."""
    pkg_dir = APPS_DIR / slug
    if not pkg_dir.exists():
        return {"error": f"no package at {pkg_dir}"}
    meta_path = pkg_dir / ".metadata.json"
    if not meta_path.exists():
        return {"error": "no metadata"}
    metadata = json.loads(meta_path.read_text())

 # Collect answers for this package (by slug+timestamp matching — Alice annotates which
 # package each answer is for via thread context; here we just take all unanswered)
    answers = _collect_answers_for_slug(slug, metadata)
    if not answers:
        return {"info": "no new answers to integrate"}

    resume_draft = (pkg_dir / "resume-draft.md").read_text()
    cover_draft = (pkg_dir / "cover-letter-draft.md").read_text()
    questions = (pkg_dir / "targeted-questions.md").read_text()

    brief = llm.load_alice_brief()
    integrate_prompt = f"""You drafted the following resume + cover letter for Jordan Avery with [FILL: ...] placeholders.
Jordan Avery has now answered some of your targeted questions. Integrate his answers and produce the FINAL versions.

RESUME DRAFT (with placeholders)
{resume_draft}

COVER LETTER DRAFT (with placeholders)
{cover_draft}

YOUR TARGETED QUESTIONS
{questions}

JORDAN AVERY'S ANSWERS
{json.dumps(answers, indent=2)}

YOUR JOB
Replace each [FILL: ...] with the substance from Jordan Avery's answer (or your best inference if his answer hints at the right detail). Preserve Jordan Avery's voice. Don't pad — keep the resume to one page. If any [FILL: ...] is still un-answerable, leave it but mark it as "[STILL NEEDED: question]".

OUTPUT FORMAT
=== RESUME FINAL ===
<the full resume markdown, placeholders replaced>

=== COVER LETTER FINAL ===
<the full cover letter markdown, placeholders replaced>
"""
    print(f"    integrating answers ({len(answers)} answers, model: {llm.MODEL_FOR_TASK['resume_draft']})...")
    res = llm.call("resume_draft", integrate_prompt, system=brief, max_tokens=3500)
    text = res["text"]

 # Split into two outputs
    parts = text.split("=== COVER LETTER FINAL ===")
    if len(parts) == 2:
        resume_final = parts[0].replace("=== RESUME FINAL ===", "").strip()
        cover_final = parts[1].strip()
    else:
        resume_final = text
        cover_final = ""

    (pkg_dir / "resume-final.md").write_text(resume_final)
    if cover_final:
        (pkg_dir / "cover-letter-final.md").write_text(cover_final)

    metadata["answers_received"] = answers
    metadata["final_generated"] = True
    metadata["final_generated_at"] = datetime.now().isoformat(timespec="seconds")
    metadata["draft_costs"]["integration"] = res["cost_usd"]
    metadata["draft_costs"]["total"] += res["cost_usd"]
    meta_path.write_text(json.dumps(metadata, indent=2))

    return {"slug": slug, "answers_integrated": len(answers), "cost": res["cost_usd"]}


def _collect_answers_for_slug(slug, metadata):
    """Read question-answers.jsonl and find answers for this package.
    Naive heuristic: take all answers received after metadata.generated_at that haven't been
    consumed yet. Improvement: track per-package answers via slug tagging in directive payload."""
    if not ANSWERS.exists():
        return []
    threshold = metadata.get("generated_at", "")
    consumed_q_nums = {a.get("q_num") for a in metadata.get("answers_received", [])}
    new = []
    with ANSWERS.open() as f:
        for line in f:
            try:
                rec = json.loads(line)
                if rec.get("received_at", "") >= threshold and rec.get("q_num") not in consumed_q_nums:
                    new.append(rec)
            except Exception:
                pass
    return new


def process_queue():
    """Pop ONE pending entry from the prep queue and generate drafts.

    Per Alice.md "If Jordan Avery has queued multiple roles (`prep order:`), Alice
    works the queue serially, one role per overnight cycle, surfacing each
    as it completes." This honors that throttle: at most one draft attempt
    per invocation. no_match entries (parse-time filter, no LLM cost) are
    flagged and skipped past — they don't consume the cycle's one attempt.
    Remaining pending entries are left untouched for the next cron cycle.
    """
    queue = safe_state.atomic_read(QUEUE, default=None)
    if queue is None:
        print("[prep: queue is empty]")
        try:
            from alice.persistence import activity_log
            activity_log.record(step="prep_materials",
                                summary="prep queue empty",
                                count=0, status="noop")
        except Exception as e:
            print(f"[activity_log: {e}]")
        return {"processed": 0}
    processed = 0
    total_cost = 0.0
    out = []
    updated_entries = {}  # substr -> updated entry; merged at end under lock
    drafted_one_this_cycle = False
    for entry in queue:
        if entry.get("status") != "pending":
            continue
        if drafted_one_this_cycle:
 # Honor one-per-cycle throttle: defer subsequent pending entries
            continue
        hit = _row_for_substr(entry.get("substr", ""))
        if not hit:
            entry["status"] = "no_match"
            updated_entries[entry.get("substr", "")] = entry
            print(f"  [prep: no unique match for {entry.get('substr')!r}; skipping]")
            continue
        row_idx, row = hit
        print(f"  [prep: generating package for {row.get('company')} - {row.get('role')}]")
        try:
            result = generate_drafts(row_idx, row, rush=entry.get("rush", False))
            entry["status"] = "drafted"
            entry["completed_at"] = datetime.now().isoformat(timespec="seconds")
            entry["slug"] = result["slug"]
            entry["cost"] = result["total_cost"]
            total_cost += result["total_cost"]
            processed += 1
            out.append(result)
        except Exception as e:
            entry["status"] = "error"
            entry["error"] = str(e)[:200]
            print(f"    ERROR: {e}")
        updated_entries[entry.get("substr", "")] = entry
        drafted_one_this_cycle = True
        deferred = sum(1 for e in queue if e.get("status") == "pending")
        if deferred > 0:
            print(f"  [prep: {deferred} pending entry/entries deferred to next cycle "
                  f"(one role per overnight cycle, per Alice.md)]")

 # Merge our updates back into the current queue state — any entries added
 # to prep-queue during our long-running LLM work are preserved.
    def merger(current):
        if current is None:
            current = []
        out_list = []
        seen_substrs = set()
        for entry in current:
            substr = entry.get("substr", "")
            if substr in updated_entries:
                out_list.append(updated_entries[substr])
            else:
                out_list.append(entry)
            seen_substrs.add(substr)
        return out_list, None

    safe_state.atomic_update(QUEUE, merger, default=[])

 # `remaining` = the queue entries we just processed; their statuses/slugs
 # were mutated in place in the loop above (drafted/error/no_match). Scanned
 # below to integrate answers for already-drafted packages.
    remaining = queue

 # Also try to integrate any pending answers for already-drafted packages
    print()
    integrations = 0
    for entry in remaining:
        if entry.get("status") == "drafted" and entry.get("slug"):
            res = integrate_answers(entry["slug"])
            if "answers_integrated" in res:
                print(f"  [integrated {res['answers_integrated']} answers for {entry['slug']}]")
                integrations += 1
                total_cost += res.get("cost", 0)

 # Aggregate sheet-write verification status across this cycle's drafts.
 # verify.verify_sheet_status_write runs per-row inside generate_drafts.
 # sheet_write_ok = API call returned. sheet_write_verified = fresh-auth
 # read-back confirmed column G shows 'materials pending'.
    sheet_writes_attempted = sum(1 for r in out if r.get("sheet_write_ok"))
    sheet_writes_verified  = sum(1 for r in out if r.get("sheet_write_verified"))
    sheet_writes_failed    = sum(1 for r in out if r.get("sheet_write_error"))
    sheet_writes_unverified = sheet_writes_attempted - sheet_writes_verified

    print()
    print(f"summary: {processed} packages drafted, {integrations} integrations, "
          f"${total_cost:.4f} spent")
    if sheet_writes_attempted or sheet_writes_failed:
        print(f"  sheet writes: {sheet_writes_attempted} attempted, "
              f"{sheet_writes_verified} VERIFIED via fresh-auth read-back, "
              f"{sheet_writes_unverified} unverified, "
              f"{sheet_writes_failed} failed")
    try:
        from alice.persistence import activity_log
        slugs = [e.get("slug") for e in remaining if e.get("status") == "drafted" and e.get("completed_at", "").startswith(datetime.now().strftime("%Y-%m-%d"))]
        summary_parts = []
        if processed:
            if sheet_writes_unverified or sheet_writes_failed:
                qual = f"sheet write attempted, {sheet_writes_verified}/{sheet_writes_attempted} verified"
            else:
                qual = "sheet write verified via fresh-auth read-back"
            summary_parts.append(
                f"{processed} package{'' if processed == 1 else 's'} drafted ({qual})"
            )
        if integrations:
            summary_parts.append(f"{integrations} answer integration{'' if integrations == 1 else 's'}")
        summary = ", ".join(summary_parts) if summary_parts else "queue processed, no new work"
        activity_log.record(
            step="prep_materials",
            summary=summary, count=processed + integrations, cost=total_cost,
            status="ok" if (processed + integrations) > 0 else "noop",
            details={
                "slugs": slugs,
                "integrations": integrations,
                "sheet_writes_attempted":  sheet_writes_attempted,
                "sheet_writes_verified":   sheet_writes_verified,
                "sheet_writes_unverified": sheet_writes_unverified,
                "sheet_writes_failed":     sheet_writes_failed,
                "verification_status":     "C2_verifier_wired",
            },
        )
    except Exception as e:
        print(f"[activity_log: {e}]")
    return {"processed": processed, "integrations": integrations, "cost": total_cost}


if __name__ == "__main__":
    process_queue()
